import io

from django.contrib.auth import get_user_model
from django.test import TestCase
from openpyxl import load_workbook

from antigenapi.models import (
    Antigen,
    Cohort,
    ElisaPlate,
    ElisaWell,
    Library,
    Llama,
    Project,
    SequencingRun,
)


class SequencingSubmissionFileTests(TestCase):
    """Regression tests: a naive library split across two ELISA plates at the
    same pan round produced duplicate sample names in the downloaded
    sequencing submission file, and never included the library name.
    """

    def setUp(self):
        user = get_user_model().objects.create(username="tester")
        antigen = Antigen.objects.create(short_name="AG1", added_by=user)
        llama = Llama.objects.create(name="Larry", added_by=user)
        cohort = Cohort.objects.create(
            cohort_num=10, llama=llama, is_naive=True, added_by=user
        )
        project = Project.objects.create(
            title="Project One", short_title="P1", added_by=user
        )
        library = Library.objects.create(project=project, cohort=cohort, added_by=user)

        # Two ELISA plates picked for the SAME (naive) library at the SAME pan
        # round concentration.
        self.plate_a = ElisaPlate.objects.create(
            library=library, pan_round_concentration=10, added_by=user
        )
        self.plate_b = ElisaPlate.objects.create(
            library=library, pan_round_concentration=10, added_by=user
        )

        # Both plates have a well at the same location ("A1") for the same
        # antigen - without disambiguation, both generate "AG1_10A1".
        ElisaWell.objects.create(plate=self.plate_a, location=1, antigen=antigen)
        ElisaWell.objects.create(plate=self.plate_b, location=1, antigen=antigen)

        self.sequencing_run = SequencingRun.objects.create(
            plate_thresholds=[],
            wells=[
                {
                    "elisa_well": {"plate": self.plate_a.pk, "location": 1},
                    "plate": 0,
                    "location": 1,  # submission plate position "A1"
                },
                {
                    "elisa_well": {"plate": self.plate_b.pk, "location": 1},
                    "plate": 0,
                    "location": 61,  # submission plate position "F1"
                },
            ],
            added_by=user,
        )

    def _sample_names(self):
        response = self.client.get(
            f"/api/sequencingrun/{self.sequencing_run.pk}/submissionfile/0/"
        )
        assert response.status_code == 200

        wb = load_workbook(io.BytesIO(response.content))
        ws = wb.worksheets[0]

        names = {}
        for row in range(3, 100):
            position = ws[f"A{row}"].value
            sample_name = ws[f"B{row}"].value
            if sample_name is not None:
                names[position] = sample_name
        return names

    def test_sample_names_are_unique_across_colliding_plates(self):
        names = self._sample_names()

        assert names["A1"] != names["F1"], (
            f"Both wells generated the same sample name: {names['A1']!r}"
        )

    def test_sample_names_include_library_cohort_suffix(self):
        names = self._sample_names()

        # Naive cohort 10 -> "_CN10" suffix.
        assert names["A1"].endswith("_CN10")
        assert names["F1"].endswith("_CN10")

    def test_sample_names_include_plate_disambiguation_suffix(self):
        names = self._sample_names()

        assert names["A1"] == "AG1_10A1.1_CN10"
        assert names["F1"] == "AG1_10A1.2_CN10"
